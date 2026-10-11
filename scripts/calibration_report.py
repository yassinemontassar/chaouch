"""Judge calibration (spec 8.2, ADR 0009): agreement between the judge and a second grader.

Usage:
  uv run tasks.py calibrate-report   # the stored v1 judge results vs the graded samples A and B
  uv run tasks.py calibrate          # re-judge the sample answers with the current judge prompts
  uv run tasks.py calibrate --cached # the same from the cache only: no judge calls
Prints a Markdown report per sample; paste the numbers into an ADR. Ungraded rows are skipped.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import sys
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

import openai
from openpyxl import load_workbook

from chaouch.config import Settings, load_models_config
from chaouch.llm import CacheMiss, LLMClient, QuotaExhausted, SQLiteCache
from evals.config import load_eval_config
from evals.gold import load_gold
from evals.judge.judge import Judge, JudgeError

DIR = Path(__file__).resolve().parents[1] / "data" / "evals" / "calibration"
SHEET_A = DIR / "judge_calibration_v1.xlsx"
JUDGE_A = DIR / "judge_calibration_v1_judge.json"
SHEET_B = DIR / "judge_calibration_v1b.xlsx"
PASS = 0.8  # facts agreement and dialect within ±1, on each sample (ADR 0009)
GEMINI = {"B2", "B2_smoke"}  # Google answers, judged by Gemma (Google): reported apart (ADR 0010)

Key = tuple[str, str]  # (question id, system)
Judged = dict[Key, dict[str, Any]]  # fact_verdicts, dialect_score, dialect_label[, fact_reasons]


@dataclass(frozen=True)
class Sample:
    name: str
    facts: dict[tuple[str, str, int], bool]  # second grader: (id, system, fact_no) -> present
    dialect: dict[Key, int]
    notes: dict[Key, str]
    answers: dict[Key, str]
    judge: Judged  # the stored v1 judge results


def rows(ws: Any) -> list[tuple[int, dict[str, Any]]]:
    values = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(values)]
    return [
        (number, dict(zip(header, row, strict=False)))
        for number, row in enumerate(values, start=2)
        if row and row[0]
    ]


def yes_no(value: Any) -> bool | None:
    text = str(value or "").strip().upper()
    return {"Y": True, "N": False}.get(text)


def score(value: Any) -> int | None:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return None
    return n if 1 <= n <= 5 else None


def add_note(notes: dict[Key, str], key: Key, text: Any) -> None:
    if text:
        notes[key] = "; ".join(filter(None, [notes.get(key), str(text)]))


def read_sheet(path: Path, judge: list[dict[str, Any]]) -> Sample:
    """Sample A: tab "Grade", one row per fact; the judge's results are in a separate file."""
    facts: dict[tuple[str, str, int], bool] = {}
    dialect: dict[Key, int] = {}
    notes: dict[Key, str] = {}
    answers: dict[Key, str] = {}
    problems = []
    for number, row in rows(load_workbook(path, read_only=True)["Grade"]):
        key = (str(row["id"]), str(row["system"]))
        if row.get("model answer"):
            answers[key] = str(row["model answer"])
        present = yes_no(row.get("owner_fact_present"))
        if present is not None:
            facts[(*key, int(row["fact_no"]))] = present
        elif row.get("owner_fact_present"):
            problems.append(f"row {number}: owner_fact_present must be Y or N")
        if row.get("owner_dialect") is not None:
            if (s := score(row["owner_dialect"])) is None:
                problems.append(f"row {number}: owner_dialect must be 1-5")
            else:
                dialect[key] = s
        add_note(notes, key, row.get("owner_notes"))
    if problems:
        raise SystemExit("Fix the sheet first:\n" + "\n".join(problems))
    by_key = {(j["id"], j["system"]): j for j in judge}
    return Sample("A", facts, dialect, notes, answers, by_key)


def graded(row: dict[str, Any], claude: Any, judge: Any) -> Any:
    """Sample B owner column: blank or "claude" keeps Claude's grade, "judge" takes the judge's,
    and on the Facts tab Y or N is the owner's own grade."""
    owner = str(row.get("owner") or "claude").strip().lower()
    return {"claude": claude, "judge": judge, "y": True, "n": False}.get(owner)


def read_sheet_b(path: Path) -> Sample:
    """Sample B: tabs "Facts" and "Dialect", Claude's and the judge's grades side by side."""
    wb = load_workbook(path, read_only=True)
    facts: dict[tuple[str, str, int], bool] = {}
    dialect: dict[Key, int] = {}
    notes: dict[Key, str] = {}
    answers: dict[Key, str] = {}
    judge: Judged = {}
    problems = []
    for number, row in rows(wb["Facts"]):
        key = (str(row["id"]), str(row["system"]))
        answers[key] = str(row["model answer"])
        claude, judged = yes_no(row["Claude"]), yes_no(row["judge"])
        present = graded(row, claude, judged)
        if claude is None or judged is None or present is None:
            problems.append(
                f"Facts row {number}: Claude/judge must be Y or N, owner claude or judge"
            )
            continue
        facts[(*key, int(row["fact #"]) + 1)] = present  # "fact #" counts from 0
        verdicts = judge.setdefault(key, {"fact_verdicts": []})["fact_verdicts"]
        verdicts.append("present" if judged else "absent")
        add_note(notes, key, row.get("notes"))
    for number, row in rows(wb["Dialect"]):
        key = (str(row["id"]), str(row["system"]))
        claude, judged = score(row["Claude (1-5)"]), score(row["judge (1-5)"])
        chosen = score(row.get("owner score (optional)")) or graded(row, claude, judged)
        if claude is None or judged is None or chosen is None:
            problems.append(f"Dialect row {number}: scores must be 1-5, owner claude or judge")
            continue
        dialect[key] = chosen
        judge.setdefault(key, {}).update(dialect_score=judged, dialect_label=None)
        add_note(notes, key, row.get("Claude reason"))
        add_note(notes, key, row.get("notes"))
    if problems:
        raise SystemExit(f"Fix {path.name} first:\n" + "\n".join(problems))
    return Sample("B", facts, dialect, notes, answers, judge)


def subset(sample: Sample, systems: set[str]) -> Sample:
    def keep(d: dict[Any, Any]) -> dict[Any, Any]:
        return {k: v for k, v in d.items() if k[1] in systems}

    return replace(
        sample,
        name=f"{sample.name}, {'/'.join(sorted(systems))} answers only",
        facts=keep(sample.facts),
        dialect=keep(sample.dialect),
        notes=keep(sample.notes),
        answers=keep(sample.answers),
    )


def cohen_kappa(pairs: list[tuple[bool, bool]]) -> float | None:
    """Cohen's kappa for two raters giving yes/no labels; None when chance agreement is total."""
    n = len(pairs)
    if not n:
        return None
    observed = sum(a == b for a, b in pairs) / n
    p_a, p_b = sum(a for a, _ in pairs) / n, sum(b for _, b in pairs) / n
    expected = p_a * p_b + (1 - p_a) * (1 - p_b)
    return None if expected == 1 else (observed - expected) / (1 - expected)


def ranks(values: list[float]) -> list[float]:
    order = sorted(range(len(values)), key=lambda i: values[i])
    out = [0.0] * len(values)
    i = 0
    while i < len(order):  # ties share the average rank
        j = i
        while j + 1 < len(order) and values[order[j + 1]] == values[order[i]]:
            j += 1
        for k in range(i, j + 1):
            out[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return out


def spearman(xs: list[float], ys: list[float]) -> float | None:
    rx, ry = ranks(xs), ranks(ys)
    n = len(xs)
    if n < 2:
        return None
    mx, my = sum(rx) / n, sum(ry) / n
    cov = sum((a - mx) * (b - my) for a, b in zip(rx, ry, strict=True))
    var = (sum((a - mx) ** 2 for a in rx) * sum((b - my) ** 2 for b in ry)) ** 0.5
    return None if var == 0 else cov / var


def pct(part: int, whole: int) -> str:
    return f"{part / whole:.0%}" if whole else "–"


def fmt(value: float | None) -> str:
    return "–" if value is None else f"{value:.2f}"


def report(sample: Sample, judge: Judged, title: str) -> tuple[str, bool]:
    """Markdown report and whether the sample passes (facts and dialect within ±1 ≥ PASS)."""
    lines = [f"# {title}: sample {sample.name}", ""]
    keys = {f[:2] for f in sample.facts} | set(sample.dialect)  # judge may cover other samples
    failed = {k: judge[k]["error"] for k in keys if judge[k].get("error")}

    fact_rows = []  # (key, fact_no, grader, judge verdict, judge reason)
    for (qid, system, n), present in sorted(sample.facts.items()):
        j = judge[(qid, system)]
        if (qid, system) not in failed:
            reason = (j.get("fact_reasons") or [""] * n)[n - 1]
            fact_rows.append(((qid, system), n, present, j["fact_verdicts"][n - 1], reason))
    pairs = [(o, v == "present") for _, _, o, v, _ in fact_rows]
    agree = sum(o == j for o, j in pairs)
    stricter = [r for r in fact_rows if r[2] and r[3] != "present"]
    looser = [r for r in fact_rows if not r[2] and r[3] == "present"]
    lines += [
        f"## Facts ({len(pairs)} graded)",
        "",
        f"- Agreement: {pct(agree, len(pairs))} ({agree}/{len(pairs)})",
        f"- Cohen's kappa: {fmt(cohen_kappa(pairs))}",
        f"- Judge stricter (grader Y, judge N): {len(stricter)}",
        f"- Judge looser (grader N, judge Y): {len(looser)}",
        "",
    ]

    dialect = [
        (k, s, judge[k]["dialect_score"])
        for k, s in sorted(sample.dialect.items())
        if k not in failed
    ]
    diffs = [abs(s - j) for _, s, j in dialect]
    within = sum(d <= 1 for d in diffs)
    lines += [
        f"## Dialect ({len(dialect)} graded)",
        "",
        f"- Mean absolute difference: {fmt(sum(diffs) / len(diffs) if diffs else None)}",
        f"- Within ±1: {pct(within, len(diffs))} ({within}/{len(diffs)})",
        f"- Spearman: {fmt(spearman([s for _, s, _ in dialect], [j for _, _, j in dialect]))}",
        "",
        "## Disagreements",
        "",
        "Every fact disagreement and every dialect difference of 2 or more.",
        "",
        "| id | system | what | grader | judge | grader notes |",
        "|---|---|---|---|---|---|",
    ]
    for key, n, present, verdict, reason in stricter + looser:
        shown = f"{verdict}: {reason}" if reason else verdict
        lines.append(
            f"| {key[0]} | {key[1]} | fact {n} | {'Y' if present else 'N'} | {shown} | "
            f"{sample.notes.get(key, '')} |"
        )
    for key, s, judged in dialect:
        if abs(s - judged) >= 2:
            label = judge[key].get("dialect_label")
            shown = f"{judged} ({label})" if label else str(judged)
            lines.append(
                f"| {key[0]} | {key[1]} | dialect | {s} | {shown} | {sample.notes.get(key, '')} |"
            )
    for key, error in sorted(failed.items()):
        lines.append(f"| {key[0]} | {key[1]} | not judged | | {error[:120]} | |")

    passed = (
        not failed
        and bool(pairs and diffs)
        and agree / len(pairs) >= PASS
        and within / len(diffs) >= PASS
    )
    return "\n".join(lines) + "\n", passed


async def rejudge(samples: list[Sample], cache_only: bool) -> tuple[Judged, list[str]]:
    """Judge every sample answer again with the current prompts (cached; overlaps judged once)."""
    cfg = load_eval_config()
    llm = LLMClient(load_models_config(), Settings(), cache_only=cache_only)
    judge = Judge(llm, cfg.judge_model, SQLiteCache(cfg.cache_path))
    gold = {q.id: q for q in load_gold()}
    answers = {k: a for s in samples for k, a in s.answers.items()}

    async def one(key: Key, answer: str) -> dict[str, Any]:
        q = gold[key[0]]
        try:
            (_, verdicts, reasons), (score_, label) = await asyncio.gather(
                judge.correctness(q.question, q.expected_answer, q.required_facts, answer),
                judge.dialect(q.question, answer),
            )
        except (QuotaExhausted, CacheMiss, JudgeError, openai.APIError) as exc:
            return {"error": f"{type(exc).__name__}: {exc}"}
        return {
            "fact_verdicts": verdicts,
            "fact_reasons": reasons,
            "dialect_score": score_,
            "dialect_label": label,
        }

    results = await asyncio.gather(*(one(k, a) for k, a in answers.items()))
    return dict(zip(answers, results, strict=True)), judge.prompt_versions


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rejudge", action="store_true", help="judge again with current prompts")
    parser.add_argument("--cached", action="store_true", help="with --rejudge: no judge calls")
    args = parser.parse_args()
    samples = [
        read_sheet(SHEET_A, json.loads(JUDGE_A.read_text(encoding="utf-8"))),
        read_sheet_b(SHEET_B),
    ]
    empty = [s.name for s in samples if not s.facts and not s.dialect]
    if empty:
        print(f"No grades in sample {', '.join(empty)} yet.")
        return 1
    if not args.rejudge:
        for s in samples:
            print(report(s, s.judge, "Judge calibration v1")[0])
        return 0

    judged, versions = asyncio.run(rejudge(samples, args.cached))
    results = [report(s, judged, f"Judge calibration ({', '.join(versions)})") for s in samples]
    for text, _ in results:
        print(text)
    for s in samples:  # same-family check, not part of the pass rule
        print(report(subset(s, GEMINI), judged, "Same-family check")[0])
    if errors := [j["error"] for j in judged.values() if j.get("error")]:
        print(f"{len(errors)} answers not judged; re-run once the quota resets. First error:")
        print(errors[0], "\n")
    for s, (_, passed) in zip(samples, results, strict=True):
        print(f"Sample {s.name}: {'PASS' if passed else 'FAIL'}")
    return 0 if all(passed for _, passed in results) else 1


if __name__ == "__main__":
    sys.exit(main())
