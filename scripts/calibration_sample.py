"""Build the blind judge-calibration sheet from already-judged answers in reports/ (no LLM calls).

Usage: uv run python scripts/calibration_sample.py
Writes data/evals/calibration/judge_calibration_v1.xlsx (for the owner) and
judge_calibration_v1_judge.json (the judge's verdicts, kept out of the sheet so grading is blind).
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUT_DIR = ROOT / "data" / "evals" / "calibration"
SHEET = OUT_DIR / "judge_calibration_v1.xlsx"
JUDGE = OUT_DIR / "judge_calibration_v1_judge.json"

SMOKE_ITEMS = ["b3-005", "birth-003"]  # owner suspects the judge was too strict on these
QUOTAS = {"B2": 8, "B3": 5, "B0_qwen": 3, "B0_falcon": 2}
SCRIPTS = ["arabic", "arabizi", "mixed_fr"]
COLUMNS = [
    "id", "system", "question", "fact_no", "gold required fact", "model answer",
    "owner_fact_present", "owner_dialect", "owner_notes",
]  # fmt: skip
GUIDE = [
    "Grade each answer without looking at the judge's scores (they are in a separate file).",
    "",
    "owner_fact_present (one per fact row):",
    "  Y = the answer states this fact, in any wording or language (Derja, Arabic, French).",
    "  N = the fact is missing, or the answer says something incompatible with it.",
    "",
    "owner_dialect (once per answer, on its first row; the grey cells below it stay empty):",
    "  5 = natural Tunisian Derja throughout",
    "  4 = mostly Derja, a few MSA or French phrases",
    "  3 = an even mix of Derja and MSA/French",
    "  2 = mostly MSA or French, a few Derja words",
    "  1 = no Derja at all",
    "  Rate the language only, not the content.",
    "",
    "owner_notes: anything useful, e.g. why a fact is borderline.",
    "",
    "These are the same rules the judge was given (evals/judge/prompts/*_v1.md).",
]


def latest(pattern: str, *, smoke: bool = False) -> Path:
    paths = sorted(p for p in REPORTS.glob(pattern) if ("_smoke_" in p.name) == smoke)
    if not paths:
        raise SystemExit(f"no report matching reports/{pattern}; run the evals first")
    return paths[-1]


def scored(report: Path) -> list[dict[str, Any]]:
    data = json.loads(report.read_text(encoding="utf-8"))
    model = data["model"]["name"]
    return [
        {**r, "model": model, "report": report.name}
        for r in data["results"]
        if r.get("fact_verdicts") and r.get("dialect_score") is not None
    ]


def spread(pool: list[dict[str, Any]], k: int, offset: int) -> list[dict[str, Any]]:
    """k items, cycling through the scripts and spaced over each script's score range."""
    ranked = sorted(pool, key=lambda r: (r["correctness"], r["dialect_score"], r["id"]))
    picks: list[dict[str, Any]] = []
    for j in range(k):
        script = SCRIPTS[(j + offset) % len(SCRIPTS)]
        left = [r for r in ranked if r not in picks]
        items = [r for r in left if r["script"] == script] or left
        picks.append(items[round(j / max(k - 1, 1) * (len(items) - 1))])
    return picks


def build_sample() -> list[dict[str, Any]]:
    smoke = {r["id"]: r for r in scored(latest("*_B2.json", smoke=True))}
    sample = [{**smoke[i], "system": "B2_smoke"} for i in SMOKE_ITEMS if i in smoke]
    for system, k in QUOTAS.items():
        sample += spread(scored(latest(f"*_{system}.json")), k, offset=len(sample))
    return sorted(sample, key=lambda r: (r["id"], r["system"]))


def write_sheet(sample: list[dict[str, Any]], gold: dict[str, dict[str, Any]]) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 100

    ws = wb.create_sheet("Grade")
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    grey = PatternFill("solid", fgColor="DDDDDD")
    wrap = Alignment(wrap_text=True, vertical="top")
    yes_no = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    one_to_five = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    ws.add_data_validation(yes_no)
    ws.add_data_validation(one_to_five)

    for item in sample:
        facts = gold[item["id"]]["required_facts"]
        first = ws.max_row + 1
        for n, fact in enumerate(facts, start=1):
            question = gold[item["id"]]["question"] if n == 1 else None
            answer = item["answer"] if n == 1 else None
            ws.append([item["id"], item["system"], question, n, fact, answer, None, None, None])
            yes_no.add(ws.cell(ws.max_row, 7))
            if n > 1:
                ws.cell(ws.max_row, 8).fill = grey
        one_to_five.add(ws.cell(first, 8))
        last = ws.max_row
        if last > first:
            for col in (3, 6):  # question and answer span the answer's fact rows
                ws.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    for col, width in zip("ABCDEFGHI", [11, 11, 40, 7, 40, 70, 12, 12, 30], strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.active = ws
    wb.save(SHEET)


def main() -> None:
    gold = {
        q["id"]: q
        for q in map(json.loads, (ROOT / "data/evals/gold/questions.jsonl").open(encoding="utf-8"))
    }
    sample = build_sample()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_sheet(sample, gold)
    judge = [
        {k: r[k] for k in ("id", "system", "model", "report", "correctness", "fact_verdicts",
                           "dialect_score", "dialect_label")}
        for r in sample
    ]  # fmt: skip
    JUDGE.write_text(json.dumps(judge, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    rows = sum(len(gold[r["id"]]["required_facts"]) for r in sample)
    print(f"{len(sample)} answers, {rows} fact rows -> {SHEET.relative_to(ROOT)}")
    for r in sample:
        print(f"  {r['id']:10} {r['system']:10} {r['script']:9} judge correctness "
              f"{r['correctness']:.2f} dialect {r['dialect_score']}")  # fmt: skip


if __name__ == "__main__":
    main()
