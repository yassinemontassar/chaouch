"""Build the native-speaker dialect sheet (ADR 0010) from already-judged answers (no LLM calls).

Usage: uv run python -m scripts.dialect_sample
Writes data/evals/calibration/dialect_check_v1.xlsx (for the owner) and dialect_check_v1_judge.json
(the judge's dialect scores, kept out of the sheet so grading is blind). Answers already in
calibration samples A and B are left out, and so are echoes (scored 1 without the judge).
"""

from __future__ import annotations

import json
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from scripts.calibration_sample import OUT_DIR, ROOT, latest, scored, spread

SHEET = OUT_DIR / "dialect_check_v1.xlsx"
JUDGE = OUT_DIR / "dialect_check_v1_judge.json"
QUOTAS = {"B1_falcon": 6, "B2": 6, "B1_qwen": 4, "B0_falcon": 2, "B0_qwen": 2}
COLUMNS = ["id", "system", "script", "question", "model answer", "owner_dialect", "owner_notes"]
GUIDE = [
    "Rate the DIALECT of each answer, 1-5. Judge the language only, not the content.",
    "The judge's scores are in a separate file: don't open it before grading.",
    "",
    "  5 = natural Tunisian Derja throughout",
    "  4 = mostly Derja, a few formal Arabic (MSA) or French phrases",
    "  3 = an even mix of Derja and MSA/French",
    "  2 = mostly MSA or French, a few Derja words",
    "  1 = no Derja at all",
    "",
    "The judge was also told (prompt dialect_v2):",
    "  - French words inside Derja sentences are normal Tunisian usage.",
    "  - Clear Moroccan, Egyptian, Gulf or Levantine words: at most 2.",
    "  - Mostly formal Arabic: at most 2. Mostly or entirely French: at most 2.",
    "  - Words copied from the question don't count as Derja.",
    "",
    "If you disagree with one of these rules, say so in owner_notes.",
]


def already_sampled() -> set[tuple[str, str]]:
    keys = set()
    a = load_workbook(OUT_DIR / "judge_calibration_v1.xlsx", read_only=True)["Grade"]
    b = load_workbook(OUT_DIR / "judge_calibration_v1b.xlsx", read_only=True)["Dialect"]
    for ws, qid, system in ((a, 0, 1), (b, 1, 0)):
        keys |= {(r[qid], r[system]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    return keys


def build_sample() -> list[dict[str, Any]]:
    used = already_sampled()
    sample: list[dict[str, Any]] = []
    for system, k in QUOTAS.items():
        pool = [
            r
            for r in scored(latest(f"*_{system}.json"))
            if (r["id"], system) not in used and not r.get("echo")
        ]
        key = lambda r: (r["dialect_score"], r["id"])  # noqa: E731
        sample += spread(pool, k, offset=len(sample), key=key)
    return sorted(sample, key=lambda r: (r["id"], r["system"]))


def write_sheet(sample: list[dict[str, Any]], gold: dict[str, dict[str, Any]]) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    ws = wb.create_sheet("Dialect")
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    one_to_five = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    ws.add_data_validation(one_to_five)
    for r in sample:
        ws.append([r["id"], r["system"], r["script"], gold[r["id"]]["question"], r["answer"]])
        one_to_five.add(ws.cell(ws.max_row, 6))
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    for col, width in zip("ABCDEFG", [11, 11, 9, 40, 90, 13, 30], strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.active = ws
    wb.save(SHEET)


def main() -> None:
    gold = {
        q["id"]: q
        for q in map(json.loads, (ROOT / "data/evals/gold/questions.jsonl").open(encoding="utf-8"))
    }
    sample = build_sample()
    write_sheet(sample, gold)
    judge = [
        {k: r[k] for k in ("id", "system", "model", "report", "dialect_score", "dialect_label")}
        for r in sample
    ]
    JUDGE.write_text(json.dumps(judge, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(sample)} answers -> {SHEET.relative_to(ROOT)}")
    for r in sample:
        print(f"  {r['id']:10} {r['system']:10} {r['script']:9} judge dialect {r['dialect_score']}")


if __name__ == "__main__":
    main()
