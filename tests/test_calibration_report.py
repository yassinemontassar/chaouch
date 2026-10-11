from pathlib import Path

import pytest
from openpyxl import Workbook

from scripts.calibration_report import (
    cohen_kappa,
    read_dialect_sheet,
    read_sheet,
    read_sheet_b,
    report,
    spearman,
    subset,
)

HEADER = [
    "id", "system", "question", "fact_no", "gold required fact", "model answer",
    "owner_fact_present", "owner_dialect", "owner_notes",
]  # fmt: skip
JUDGE = [
    {"id": "q1", "system": "B2", "fact_verdicts": ["present", "absent"], "dialect_score": 4,
     "dialect_label": "derja"},
    {"id": "q2", "system": "B3", "fact_verdicts": ["present"], "dialect_score": 1,
     "dialect_label": "msa"},
]  # fmt: skip


def sheet(tmp_path: Path, rows: list[list[object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Grade"
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    path = tmp_path / "grades.xlsx"
    wb.save(path)
    return path


def test_kappa_and_spearman() -> None:
    assert cohen_kappa([(True, True), (False, False)]) == 1.0
    assert cohen_kappa([(True, False), (False, True)]) == -1.0
    assert cohen_kappa([(True, True), (True, True)]) is None  # no variation
    assert spearman([1, 2, 3], [10, 20, 30]) == pytest.approx(1.0)
    assert spearman([1, 2, 3], [3, 2, 1]) == pytest.approx(-1.0)
    assert spearman([1, 1, 1], [1, 2, 3]) is None


def test_report_counts_direction_and_disagreements(tmp_path: Path) -> None:
    path = sheet(
        tmp_path,
        [
            ["q1", "B2", "Q?", 1, "fact a", "answer", "y", 4, None],
            ["q1", "B2", None, 2, "fact b", None, "Y", None, "paraphrased"],  # judge said absent
            ["q2", "B3", "Q?", 1, "fact c", "answer", "N", 4, None],  # judge said present
        ],
    )
    sample = read_sheet(path, JUDGE)
    text, passed = report(sample, sample.judge, "v1")
    assert not passed
    assert "Agreement: 33% (1/3)" in text
    assert "Judge stricter (grader Y, judge N): 1" in text
    assert "Judge looser (grader N, judge Y): 1" in text
    assert "Within ±1: 50% (1/2)" in text
    assert "| q1 | B2 | fact 2 | Y | absent | paraphrased |" in text
    assert "| q2 | B3 | dialect | 4 | 1 (msa) |" in text


def test_blank_rows_are_skipped_and_bad_values_rejected(tmp_path: Path) -> None:
    owner = read_sheet(
        sheet(tmp_path, [["q1", "B2", "Q?", 1, "fact a", "answer", None, None, None]]), JUDGE
    )
    assert not owner.facts
    assert not owner.dialect
    assert owner.answers == {("q1", "B2"): "answer"}
    with pytest.raises(SystemExit, match="row 2: owner_fact_present"):
        read_sheet(sheet(tmp_path, [["q1", "B2", "Q?", 1, "f", "a", "maybe", None, None]]), JUDGE)
    with pytest.raises(SystemExit, match="row 2: owner_dialect"):
        read_sheet(sheet(tmp_path, [["q1", "B2", "Q?", 1, "f", "a", "Y", 7, None]]), JUDGE)


def sheet_b(tmp_path: Path, owner: str | None = None) -> Path:
    wb = Workbook()
    facts = wb.active
    facts.title = "Facts"
    facts.append(["system", "id", "script", "question", "fact #", "required fact", "model answer",
                  "Claude", "judge", "agree", "owner", "notes"])  # fmt: skip
    facts.append(["B0", "q1", "arabic", "Q?", 0, "fact a", "answer", "Y", "N", "NO", owner, None])
    facts.append(["B0", "q1", "arabic", "Q?", 1, "fact b", "answer", "N", "N", "yes", None, None])
    dialect = wb.create_sheet("Dialect")
    dialect.append(["system", "id", "script", "question", "model answer", "Claude (1-5)",
                    "judge (1-5)", "gap", "Claude reason", "owner", "owner score (optional)",
                    "notes"])  # fmt: skip
    dialect.append(["B0", "q1", "arabic", "Q?", "answer", 1, 5, 4, "Moroccan", owner, None, None])
    path = tmp_path / "b.xlsx"
    wb.save(path)
    return path


def test_sample_b_counts_facts_from_zero_and_takes_owner_picks(tmp_path: Path) -> None:
    sample = read_sheet_b(sheet_b(tmp_path))
    assert sample.facts == {("q1", "B0", 1): True, ("q1", "B0", 2): False}
    assert sample.judge[("q1", "B0")] == {
        "fact_verdicts": ["absent", "absent"], "dialect_score": 5, "dialect_label": None,
    }  # fmt: skip
    text, passed = report(sample, sample.judge, "v1")
    assert "| q1 | B0 | dialect | 1 | 5 | Moroccan |" in text
    assert not passed

    judge_wins = read_sheet_b(sheet_b(tmp_path, owner="judge"))
    assert judge_wins.facts[("q1", "B0", 1)] is False
    assert judge_wins.dialect[("q1", "B0")] == 5
    with pytest.raises(SystemExit, match="Facts row 2"):
        read_sheet_b(sheet_b(tmp_path, owner="maybe"))


def test_rejudged_reasons_errors_and_pass(tmp_path: Path) -> None:
    path = sheet(tmp_path, [["q1", "B2", "Q?", 1, "fact a", "answer", "Y", 4, None]])
    sample = read_sheet(path, JUDGE)
    fresh = {("q1", "B2"): {"fact_verdicts": ["absent"], "fact_reasons": ["no price"],
                            "dialect_score": 5, "dialect_label": "derja"}}  # fmt: skip
    text, passed = report(sample, fresh, "v2")
    assert "| q1 | B2 | fact 1 | Y | absent: no price |" in text
    assert not passed  # 0% facts
    fresh[("q1", "B2")]["fact_verdicts"] = ["present"]
    assert report(sample, fresh, "v2")[1]
    other_sample = {("q9", "B0"): {"error": "quota: groq"}}
    assert report(sample, {**fresh, **other_sample}, "v2")[1]  # another sample's error
    text, passed = report(sample, {("q1", "B2"): {"error": "quota: groq"}}, "v2")
    assert not passed
    assert "| q1 | B2 | not judged | | quota: groq | |" in text


def test_subset_keeps_only_the_given_systems(tmp_path: Path) -> None:
    sample = read_sheet(sheet(tmp_path, [
        ["q1", "B2", "Q?", 1, "fact a", "answer", "Y", 4, None],
        ["q2", "B3", "Q?", 1, "fact c", "answer", "N", 1, None],
    ]), JUDGE)  # fmt: skip
    only = subset(sample, {"B2"})
    assert list(only.facts) == [("q1", "B2", 1)]
    assert list(only.dialect) == [("q1", "B2")]
    assert "Agreement: 100% (1/1)" in report(only, sample.judge, "v1")[0]


def test_native_dialect_sheet(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dialect"
    ws.append(
        ["id", "system", "script", "question", "model answer", "owner_dialect", "owner_notes"]
    )
    ws.append(["q1", "B2", "arabic", "Q?", "answer", 2, "Moroccan"])
    ws.append(["q2", "B3", "arabic", "Q?", "answer", None, None])  # not graded yet
    path = tmp_path / "c.xlsx"
    wb.save(path)
    sample = read_dialect_sheet(path, JUDGE)
    assert sample.dialect == {("q1", "B2"): 2}
    text = report(sample, sample.judge, "native")[0]
    assert "## Facts" not in text
    assert "Within ±1: 0% (0/1)" in text
    assert "| q1 | B2 | dialect | 2 | 4 (derja) | Moroccan |" in text
